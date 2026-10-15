"""Export saved, approved attendance reconciliation; no database/provider access."""
import json
from pathlib import Path
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

root = Path("generated-artifacts/bnms-zoom-reconciliation")
read = lambda name: json.loads((root / name).read_text())
review = read("review.json")
partial = {r["name"]: r for r in read("partial-name-review.json")}
approved = {}
for manifest, evidence in [
    ("approved-manifest.json", "apply-evidence.json"),
    ("second-approved-manifest.json", "second-apply-evidence.json"),
    ("third-approved-manifest.json", "third-apply-evidence.json"),
]:
    assert read(evidence)["committed"]
    for row in read(manifest)["approved"]:
        assert row["name"] not in approved
        approved[row["name"]] = row
assert len(approved) == 68
# Explicitly accepted alias of an already-attended booking; not a new attendee.
assert read("second-approved-manifest.json")["already_attended_unchanged"] == "Ruth Cornwell"
approved["Ruth"] = {**approved["Ruth Cornwell"],
    "basis": "Approved alias of Ruth Cornwell; booking already marked attended"}

wb = Workbook()
unmatched = wb.active
unmatched.title = "Unmatched and ambiguous"
waiting = wb.create_sheet("Waiting room only")
matched = wb.create_sheet("Approved matches")
summary = wb.create_sheet("Read me")
headers = ["Zoom display name", "Zoom email", "In-meeting sessions", "Reported minutes",
           "Reconciliation status", "Registrant name(s)", "Registrant email(s)",
           "Booking ID(s)", "Match basis / review notes"]
for sheet in [unmatched, waiting, matched]:
    sheet.append(headers)

def person(b):
    return " ".join(filter(None, [b.get("attendee_first_name") or b.get("first_name"),
                                 b.get("attendee_last_name") or b.get("last_name")]))

for row in review:
    name = row["name"]
    candidates, notes = [], []
    if name in approved:
        a = approved[name]
        sheet, status = matched, "Approved match"
        candidates = [a["booking"]]
        notes = [a.get("basis", "Unique exact-name match approved")]
    elif row["status"] == "Waiting room only":
        sheet, status = waiting, "No in-meeting attendance in supplied export"
        notes = ["Waiting-room record only; not counted as an attended session."]
    else:
        sheet = unmatched
        status = "Ambiguous: multiple bookings" if row["matches"] else "Unmatched"
        candidates = list(row["matches"])
        for c in partial.get(name, {}).get("candidates", []):
            if not any(b["id"] == c["booking"]["id"] for b in candidates):
                candidates.append(c["booking"])
            notes.append(c.get("basis", "Possible candidate"))
        notes.append("Candidates are not confirmed matches." if candidates else
                     "No approved registrant match found in the saved reconciliation.")
    sheet.append([name, row.get("email", ""), row["sessions"], row["minutes"], status,
                  "\n".join(person(b) for b in candidates),
                  "\n".join(b.get("attendee_email") or b.get("email") or "" for b in candidates),
                  "\n".join(b["id"] for b in candidates), "; ".join(notes)])

assert unmatched.max_row - 1 == 27
assert waiting.max_row - 1 == 3
assert matched.max_row - 1 == 69
assert sum(s.max_row - 1 for s in [unmatched, waiting, matched]) == len(review) == 99
summary.append(["Item", "Details"])
for row in [
    ("Event", "ASNC & BNMS Joint Webinar Series – ‘Quantify to Clarify’: The Power of Myocardial Blood Flow"),
    ("Event ID", "2fd9cec7-20e7-46b7-9e91-e7a62cdfad48"),
    ("Report prepared", "15/10/2026"),
    ("Evidence", "Saved Zoom reconciliation and committed administrator-approved matching results from 14/10/2026. Not a fresh live attendance audit."),
    ("Unmatched in-meeting display names", 25),
    ("Ambiguous display names with multiple bookings", 2),
    ("Waiting-room-only display names", 3),
    ("Approved matched display names", 69),
    ("Distinct bookings marked attended by reconciliation", 68),
    ("Why matched names exceed bookings", "Ruth and Ruth Cornwell were approved as the same registrant; the existing attendance was left unchanged."),
    ("Interpretation", "Zoom display names are not necessarily distinct people; devices, aliases, hosts and repeated connections may appear."),
    ("Minutes", "Reported in-meeting minutes from the saved Zoom review; not a CPD award or certificate entitlement."),
    ("Privacy", "Contains attendee personal data. Share only with authorised event administrators."),
    ("Changes made for this report", "None. No attendance, CPD points or certificates were changed."),
]:
    summary.append(row)

for sheet in wb:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="17365D")
    widths = [33, 38, 16, 18, 40, 36, 43, 42, 65] if sheet != summary else [48, 110]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Preserve source text, never executable formulas.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F4F8")
        sheet.row_dimensions[row[0].row].height = 54
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
out = root / "BNMS_Webinar_Unmatched_Attendees_15-10-2026.xlsx"
wb.save(out)
check = load_workbook(out, read_only=True)
assert check.sheetnames[0] == "Unmatched and ambiguous"
assert check.worksheets[0].max_row == 28
assert not any(c.data_type == "f" for s in check for row in s for c in row)
print(out)
print("25 unmatched; 2 ambiguous; 3 waiting-room only; 69 matched display names / 68 bookings")
