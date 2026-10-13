"""Build a confidential, formula-free review workbook from the read-only snapshot."""
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

data = json.loads(Path('/tmp/bnms-nmc-34-reconciliation.json').read_text())
rows = data['rows']
out = Path('generated-artifacts/BNMS_NMC_34_Member_Reconciliation_13-10-2026.xlsx')
out.parent.mkdir(parents=True, exist_ok=True)
wb = Workbook()
summary = wb.active
summary.title = 'Read me'

def date(s):
    return datetime.fromisoformat(s[:10]) if s else None

def add_table(name, headers, values, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in values:
        ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Never interpret imported text as an Excel formula.
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, datetime):
                cell.number_format = 'DD/MM/YYYY'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243C48')
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'C2'
    ws.row_dimensions[1].height = 34
    for i in range(2, ws.max_row+1):
        ws.row_dimensions[i].height = 95
    table = Table(displayName='Table'+str(len(wb.worksheets)), ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    return ws

def assessment(r):
    group = r['group']
    if group == 'Expired within 90 days':
        return ('No membership history. Retained expiry is within the NMC report’s 90-day allowance.',
                'BNMS to confirm the prior term was paid and the expiry is correct. If approved, establish expired history; do not mark renewed.')
    if group == 'Expired over 90 days':
        return ('No membership history. Retained expiry is beyond the NMC report’s 90-day allowance.',
                'BNMS to check for a later renewal or wrong expiry. Historical paid confirmation alone will not restore NMC eligibility.')
    if group == 'Direct Debit review':
        if not r['discoveries']:
            return ('Legacy DD flag only; no expiry, agreement, plan or payment found in checked records.',
                    'BNMS to provide membership term and payment/mandate reference. Verify identity and payment before establishing history.')
        status = sorted(r['discoveries'], key=lambda d:d['checked'])[-1]['status']
        if status == 'cancelled':
            return ('September discovery recorded a cancelled mandate. No linked payment, plan or agreement found.',
                    'Check current provider evidence and the last paid term. Cancellation does not erase a previously paid entitlement; do not restart collection.')
        return ('September discovery recorded an active mandate. No linked payment, plan or agreement found.',
                'Verify current mandate and payments in GoCardless, then confirm the paid-through date. An active mandate alone is not payment evidence.')
    if r['payments']:
        return ('Live payment record is submitted, not confirmed paid. First collection is scheduled for 15/10/2026.',
                'Await the existing collection outcome and normal reconciliation. Do not submit another collection or manually mark paid.')
    return ('Live agreement is mandate_pending; history is pending_payment_setup. No payment plan or collection recorded.',
            'Confirm whether the member completed mandate setup. Check provider completion before deciding whether setup needs resuming.')

all_values = []
detail_values = []
for r in rows:
    finding, action = assessment(r)
    h = r['history']
    all_values.append([r['name'],r['email'],r['memberClass'],r['group'],r['customStatus'],
        date(r['expiry']), max(0,r['daysExpired']) if r['daysExpired'] is not None else None,
        r['originalExpiry'], '; '.join(f"{x['status']} / {x['paymentStatus']}" for x in h) or 'No history',
        date(h[0]['end']) if h else None, finding, action, 'Awaiting review', '', r['memberId']])
    if 'review' in r['group']:
        # Discovery runs can contain repeated snapshots of the same mandate.
        unique = {}
        for d in sorted(r['discoveries'],key=lambda x:x['checked']):
            unique[d['mandate']] = d
        detail_values.append([r['name'],r['email'],r['group'],
            '; '.join(f"{a['provider']} / {a['environment']} / {a['status']}" for a in r['agreements']) or 'None recorded',
            '; '.join(f"{p['environment']} / {p['status']}; last payment: {p['lastPaymentStatus']}; instalments paid: {p['instalmentsPaid']}" for p in r['plans']) or 'None recorded',
            '\n'.join(f"{d['mandate']} — {d['status']} ({d['environment']}); saved {d['checked'][:10]}" for d in unique.values()) or 'No discovery record',
            '\n'.join(f"{p['reference']} — {p['status']} ({p['environment']}); {p['amountMinor']/100:.2f} {p['currency']}; charge {p['date']}; updated {p['updated']}" for p in r['payments']) or 'None found in checked payment records',
            r['historicalPaymentCount'],r['legacyDD'] or 'Not set', finding, action, r['memberId']])

counts = Counter(r['group'] for r in rows)
for row in [
    ['BNMS — NMC membership reconciliation', ''],
    ['Confidential', 'Contains member names and email addresses. Share only with authorised BNMS administrators.'],
    ['Report date', date(data['reportDate'])],
    ['Database snapshot (UTC)', data['asOf']],
    ['Scope', '34 members remaining after the separately approved 72-member repair. Full with NMC and Full junior with NMC classes only.'],
    ['Expired within 90 days', counts['Expired within 90 days']],
    ['Expired over 90 days', counts['Expired over 90 days']],
    ['Legacy Direct Debit review', counts['Direct Debit review']],
    ['Newer payment/setup review', counts['Payment or setup review']],
    ['Total members', len(rows)],
    ['Evidence boundary', 'Read-only production database review. No direct GoCardless or Xero API verification was performed. Absence of local payment evidence does not prove no payment occurred.'],
    ['Important payment finding', 'Two newer memberships each have a live £10.66 payment submitted for 15/10/2026. These are pending collections, not confirmed failures or paid memberships. The third awaits mandate setup.'],
    ['Legacy DD evidence', 'Three mandates last recorded active and one cancelled in September; two members have only a legacy DD flag. Repeated discovery snapshots are not separate mandates.'],
    ['Date interpretation', 'Original expiry text is retained. Parsed expiry uses the existing import contract: four-digit years DD/MM/YYYY; two-digit years M/D/YY. Review source evidence before correcting ambiguous values.'],
    ['How to use', 'Filter the Member review sheet by review group. Record decisions and supporting evidence in the final review columns. Payment evidence contains details for all nine payment-related cases.'],
    ['90-day allowance', 'This is the NMC report inclusion rule, not permission to renew, collect payment or extend account access.'],
    ['Changes made', 'None to database records, memberships, billing or providers. No migrations required.'],
]:
    summary.append(row)
summary.column_dimensions['A'].width=33
summary.column_dimensions['B'].width=110
for row in summary:
    for c in row:
        c.alignment=Alignment(vertical='top',wrap_text=True)
        if isinstance(c.value,str): c.data_type='s'
        if isinstance(c.value,datetime): c.number_format='DD/MM/YYYY'
    row[0].font=Font(bold=True,color='243C48')
    summary.row_dimensions[row[0].row].height=48
summary['A1'].font=Font(bold=True,size=16,color='243C48')
summary.freeze_panes='B2'
add_table('Member review',
    ['Member name','Email address','Member class','Review group','Custom membership status','Retained expiry',
     'Days past expiry','Original expiry text','History status / payment','History term end','Finding',
     'Recommended action / confirmation needed','Review decision','BNMS notes / evidence','Member UUID'],
    all_values,[26,34,25,26,24,16,14,20,32,16,55,65,24,45,38])
add_table('Payment evidence',
    ['Member name','Email address','Review group','Agreement evidence','Plan evidence','Saved mandate discovery',
     'Stored collection evidence','Historical DD payment count','Legacy DD flag','Finding','Recommended action','Member UUID'],
    detail_values,[26,34,26,38,45,58,70,18,18,55,65,38])
wb.save(out)
check=load_workbook(out)
assert check['Member review'].max_row == 35
assert check['Payment evidence'].max_row == 10
assert len(set(row[-1].value for row in list(check['Member review'].rows)[1:]))==34
assert not any(c.data_type=='f' for ws in check for row in ws for c in row)
print(json.dumps({'file':str(out),'members':34,'paymentReviews':9,'validated':True}))
