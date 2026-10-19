"""Read only the authorised sheet; emit exact, validated metadata, never contact Xero."""
import sys, json, hashlib, datetime
from decimal import Decimal
import openpyxl

path = sys.argv[1]
with open(path, "rb") as source:
    digest = hashlib.sha256(source.read()).hexdigest()
book = openpyxl.load_workbook(path, read_only=True, data_only=True)
worksheet = sys.argv[2] if len(sys.argv) > 2 else "To import good"
schemas = {
    "To import good": ["iConnect member UUID", "Invoice number", "Invoice date", "Gross", "Invoice description"],
    "Import foreign": ["iConnect member UUID", "Invoice number", "Invoice date", "Xero contact name", "Gross", "iConnect member name", "Invoice description"],
}
assert worksheet in schemas, "Unauthorised worksheet"
sheet = book[worksheet]
expected = schemas[worksheet]
with_status = "--status" in sys.argv[3:]
if with_status:
    expected = expected + ["Payment status"]
assert list(next(sheet.iter_rows(values_only=True))) == expected, "Unexpected worksheet headers"
rows = []
for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
    if not any(value is not None for value in row):
        continue
    values = dict(zip(expected, row))
    member, number, date, gross, description = (
        values[key] for key in ["iConnect member UUID", "Invoice number", "Invoice date", "Gross", "Invoice description"]
    )
    assert isinstance(date, (datetime.date, datetime.datetime)), f"Invalid date at row {index}"
    amount = Decimal(str(gross).replace(",", "").replace("£", "").strip()) * 100
    assert amount.is_finite() and amount == amount.to_integral_value(), f"Invalid amount at row {index}"
    assert abs(amount) < 9007199254740991, f"Amount too large at row {index}"
    number = str(number).strip() if number is not None else ""
    assert number, f"Missing reference at row {index}"
    status = {}
    if with_status:
        raw_status = values["Payment status"]
        if raw_status not in ("Paid", "Outstanding"):
            raise ValueError(f"Invalid Payment status at row {index}")
        status = dict(status=raw_status.lower(), source_status=raw_status)
    rows.append(dict(source_row=index, member_id=str(member).strip(),
                     invoice_number=number, invoice_date=date.strftime("%Y-%m-%d"),
                     amount_minor=int(amount), description=str(description or "").strip(), **status))
print(json.dumps(dict(source_sha256=digest, worksheet=sheet.title, rows=rows)))
